"""Excel / CSV / JSON exports (PRD 7.5, appendix 22.3).

The Excel file is meant to be read by people: a title block, a styled Excel
table per sheet (banded rows, built-in filters), real dates, rupiah formats,
readable column widths and a totals row. CSV and JSON keep a flat,
machine-friendly layout.
"""

from __future__ import annotations

import csv
import io
import json
import zipfile
from collections.abc import Callable
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

TYPE_LABEL = {"receipt": "Nota", "invoice": "Faktur", "other": "Lainnya"}
STATUS_LABEL = {
    "auto_approved": "Otomatis",
    "reviewed": "Ditinjau",
    "needs_review": "Perlu tinjau",
    "exported": "Diekspor",
}
PAYMENT_LABEL = {
    "cash": "Tunai",
    "debit_card": "Kartu debit",
    "credit_card": "Kartu kredit",
    "qris": "QRIS",
    "bank_transfer": "Transfer bank",
    "e_wallet": "Dompet digital",
}


@dataclass
class ExportRow:
    document_id: str
    doc_type: str
    status: str
    values: dict[str, Any]
    archive_path: str = ""
    original_name: str = ""


# CSV/JSON keep this stable layout so other tools can rely on it.
SUMMARY_COLUMNS = [
    ("Tanggal", "issue_date"), ("Vendor", "vendor_name"), ("NPWP", "vendor_tax_id"), ("Nomor dokumen", "document_number"),
    ("Jenis", "_type"), ("Jatuh tempo", "due_date"), ("Mata uang", "currency"), ("Subtotal", "subtotal"),
    ("Diskon", "discount"), ("Pajak", "tax_amount"), ("Total", "total"), ("Metode bayar", "payment_method"),
    ("Status", "_status"), ("Arsip", "_archive"), ("Berkas asli", "_original"), ("ID dokumen", "_id"),
]  # fmt: skip


def _cell(row: ExportRow, key: str) -> Any:
    special = {
        "_type": TYPE_LABEL.get(row.doc_type, row.doc_type),
        "_status": STATUS_LABEL.get(row.status, row.status),
        "_archive": row.archive_path,
        "_original": row.original_name,
        "_id": row.document_id,
    }
    return special[key] if key in special else row.values.get(key)


def _csv_safe(v: Any) -> Any:
    """Neutralise spreadsheet formula injection (=, +, -, @) in text cells."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


# --------------------------------------------------------------------------- Excel

BRAND = "1F3A5F"
MUTED = "6F6D68"
LINE = "C3C2B7"
DATE_FMT = "dd/mm/yyyy"
MONEY_FMT = '"Rp" #,##0;-"Rp" #,##0'
MONEY_FMT_DEC = '"Rp" #,##0.00;-"Rp" #,##0.00'
HEADER_ROW, FIRST_ROW = 4, 5


@dataclass
class _Col:
    header: str
    kind: str  # text | wrap | date | money | int | qty
    width: float
    get: Callable[[Any, int], Any]


def _to_date(v: Any) -> date | None:
    try:
        return date.fromisoformat(str(v)) if v else None
    except ValueError:
        return None


def _num(v: Any) -> float | None:
    try:
        return float(v) if v is not None and v != "" else None
    except (TypeError, ValueError):
        return None


def _npwp(v: Any) -> str | None:
    if not v:
        return None
    digits = "".join(ch for ch in str(v) if ch.isdigit())
    if len(digits) == 15:
        return f"{digits[:2]}.{digits[2:5]}.{digits[5:8]}.{digits[8]}-{digits[9:12]}.{digits[12:]}"
    return digits or str(v)


def _value(key: str) -> Callable[[ExportRow, int], Any]:
    return lambda r, _i: r.values.get(key)


def _summary_columns(rows: list[ExportRow]) -> list[_Col]:
    cols = [
        _Col("No", "int", 5, lambda _r, i: i),
        _Col("Tanggal", "date", 12, lambda r, _i: _to_date(r.values.get("issue_date"))),
        _Col("Jenis", "text", 9, lambda r, _i: TYPE_LABEL.get(r.doc_type, r.doc_type)),
        _Col("Vendor", "wrap", 30, _value("vendor_name")),
        _Col("NPWP", "text", 21, lambda r, _i: _npwp(r.values.get("vendor_tax_id"))),
        _Col("Nomor dokumen", "text", 18, _value("document_number")),
        _Col("Jatuh tempo", "date", 12, lambda r, _i: _to_date(r.values.get("due_date"))),
        _Col("Subtotal", "money", 15, lambda r, _i: _num(r.values.get("subtotal"))),
        _Col("Diskon", "money", 13, lambda r, _i: _num(r.values.get("discount"))),
        _Col("Pajak (PPN)", "money", 14, lambda r, _i: _num(r.values.get("tax_amount"))),
        _Col("Total", "money", 16, lambda r, _i: _num(r.values.get("total"))),
        _Col("Metode bayar", "text", 14, lambda r, _i: PAYMENT_LABEL.get(str(r.values.get("payment_method")), r.values.get("payment_method"))),
        _Col("Jumlah item", "int", 9, lambda r, _i: len(r.values.get("line_items") or [])),
        _Col("Status", "text", 13, lambda r, _i: STATUS_LABEL.get(r.status, r.status)),
        _Col("Berkas asli", "text", 22, lambda r, _i: r.original_name or None),
        _Col("Nama arsip", "text", 30, lambda r, _i: r.archive_path or None),
        _Col("ID dokumen", "text", 24, lambda r, _i: r.document_id),
    ]  # fmt: skip
    # Only show the currency when it is not all rupiah (the money format already says "Rp").
    if any((r.values.get("currency") or "IDR") != "IDR" for r in rows):
        cols.insert(7, _Col("Mata uang", "text", 9, lambda r, _i: r.values.get("currency") or "IDR"))
    return cols


ItemRec = tuple[ExportRow, dict[str, Any]]


def _item_columns() -> list[_Col]:
    return [
        _Col("No", "int", 5, lambda _x, i: i),
        _Col("Tanggal", "date", 12, lambda x, _i: _to_date(x[0].values.get("issue_date"))),
        _Col("Vendor", "wrap", 28, lambda x, _i: x[0].values.get("vendor_name")),
        _Col("Nomor dokumen", "text", 18, lambda x, _i: x[0].values.get("document_number")),
        _Col("Deskripsi", "wrap", 38, lambda x, _i: x[1].get("description")),
        _Col("Qty", "qty", 7, lambda x, _i: _num(x[1].get("qty"))),
        _Col("Harga satuan", "money", 15, lambda x, _i: _num(x[1].get("unit_price"))),
        _Col("Jumlah", "money", 16, lambda x, _i: _num(x[1].get("amount"))),
        _Col("ID dokumen", "text", 24, lambda x, _i: x[0].document_id),
    ]


def _write_sheet(
    ws: Any, title: str, subtitle: str, cols: list[_Col], records: list[Any], table_name: str
) -> None:
    ws["A1"] = title
    ws["A1"].font = Font(size=16, bold=True, color=BRAND)
    ws["A2"] = subtitle
    ws["A2"].font = Font(size=10, color=MUTED)
    ws.row_dimensions[1].height = 24

    for c, col in enumerate(cols, 1):
        cell = ws.cell(row=HEADER_ROW, column=c, value=col.header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=BRAND)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[HEADER_ROW].height = 32

    widths = [col.width for col in cols]
    for i, rec in enumerate(records, 1):
        r = FIRST_ROW + i - 1
        for c, col in enumerate(cols, 1):
            value = col.get(rec, i)
            cell = ws.cell(row=r, column=c, value=value)
            if isinstance(value, str) and value[:1] == "=":
                cell.data_type = "s"  # text that looks like a formula stays text (formula injection)
            if col.kind == "date":
                cell.number_format = DATE_FMT
                cell.alignment = Alignment(horizontal="center", vertical="top")
            elif col.kind == "money":
                cell.number_format = MONEY_FMT_DEC if value and float(value) % 1 else MONEY_FMT
                cell.alignment = Alignment(horizontal="right", vertical="top")
            elif col.kind in ("int", "qty"):
                cell.number_format = "0" if col.kind == "int" else "#,##0.##"
                cell.alignment = Alignment(horizontal="center", vertical="top")
            else:
                cell.alignment = Alignment(vertical="top", wrap_text=col.kind == "wrap")
            if col.kind == "text" and value:
                widths[c - 1] = max(widths[c - 1], min(48, len(str(value)) + 2))
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w

    if records:
        last = FIRST_ROW + len(records) - 1
        table = Table(displayName=table_name, ref=f"A{HEADER_ROW}:{get_column_letter(len(cols))}{last}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
        # Totals sit below the table as plain SUM ranges, which every spreadsheet app understands.
        tr = last + 1
        money_cols = [c for c, col in enumerate(cols, 1) if col.kind == "money"]
        rule = Border(top=Side(style="thin", color=LINE), bottom=Side(style="double", color=BRAND))
        for c in range(1, len(cols) + 1):
            ws.cell(row=tr, column=c).border = rule
        if money_cols:
            label = ws.cell(row=tr, column=min(money_cols) - 1, value="Total")
            label.font = Font(bold=True)
            label.alignment = Alignment(horizontal="right")
        for c in money_cols:
            letter = get_column_letter(c)
            cell = ws.cell(row=tr, column=c, value=f"=SUM({letter}{FIRST_ROW}:{letter}{last})")
            cell.font = Font(bold=True)
            cell.number_format = MONEY_FMT
            cell.alignment = Alignment(horizontal="right")
    else:
        ws.cell(row=FIRST_ROW, column=1, value="Tidak ada data.").font = Font(italic=True, color=MUTED)

    ws.freeze_panes = ws.cell(row=FIRST_ROW, column=1)
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"


def _ignore_number_as_text(xlsx: bytes) -> bytes:
    """Hide Excel's green 'number stored as text' marks: NPWP and document numbers are text on purpose."""
    tag = '<ignoredErrors><ignoredError sqref="A1:Z10000" numberStoredAsText="1"/></ignoredErrors>'
    src = zipfile.ZipFile(io.BytesIO(xlsx))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename.startswith("xl/worksheets/sheet") and item.filename.endswith(".xml"):
                xml = data.decode("utf-8")
                # Schema order: ignoredErrors precedes drawing, legacyDrawing, tableParts and extLst.
                for anchor in ("<drawing", "<legacyDrawing", "<tableParts", "<extLst", "</worksheet>"):
                    if anchor in xml:
                        xml = xml.replace(anchor, tag + anchor, 1)
                        break
                data = xml.encode("utf-8")
            dst.writestr(item, data)
    return out.getvalue()


def _rupiah(v: float) -> str:
    return "Rp " + f"{v:,.0f}".replace(",", ".")


def to_xlsx(rows: list[ExportRow], exported_at: datetime | None = None) -> bytes:
    when = exported_at or datetime.now()
    wb = Workbook()
    summary = wb.active
    summary.title = "Ringkasan"
    total = sum(_num(r.values.get("total")) or 0.0 for r in rows)
    _write_sheet(
        summary,
        "Rekap Dokumen",
        f"Diekspor dari DokuData pada {when:%d/%m/%Y %H:%M} · {len(rows)} dokumen · total {_rupiah(total)}",
        _summary_columns(rows),
        rows,
        "Ringkasan",
    )
    items: list[ItemRec] = [(r, li) for r in rows for li in r.values.get("line_items") or []]
    _write_sheet(
        wb.create_sheet("Item"),
        "Rincian Item",
        f"{len(items)} item dari {len(rows)} dokumen",
        _item_columns(),
        items,
        "Item",
    )
    wb.properties.title = "Rekap Dokumen DokuData"
    wb.properties.creator = "DokuData"
    buf = io.BytesIO()
    wb.save(buf)
    return _ignore_number_as_text(buf.getvalue())


# --------------------------------------------------------------------------- CSV / JSON


def to_csv(rows: list[ExportRow]) -> bytes:
    """Summary rows as CSV (UTF-8 with BOM so Excel opens it with the right encoding)."""
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow([c for c, _ in SUMMARY_COLUMNS] + ["Jumlah item"])
    for r in rows:
        w.writerow(
            [_csv_safe(_cell(r, k)) for _, k in SUMMARY_COLUMNS] + [len(r.values.get("line_items") or [])]
        )
    return ("﻿" + buf.getvalue()).encode("utf-8")


def to_json(rows: list[ExportRow]) -> bytes:
    data = [
        {"document_id": r.document_id, "document_type": r.doc_type, "status": r.status, "archive_path": r.archive_path,
         "original_name": r.original_name, **r.values}
        for r in rows
    ]  # fmt: skip
    return json.dumps({"documents": data}, ensure_ascii=False, indent=2).encode("utf-8")


EXPORTERS = {"xlsx": to_xlsx, "csv": to_csv, "json": to_json}
MIME = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "csv": "text/csv; charset=utf-8",
    "json": "application/json",
    "zip": "application/zip",
}
