import csv
import io
import json

import pytest
from openpyxl import load_workbook

from pipeline.archive import archive_path, dedupe_paths, validate_pattern
from pipeline.export import ExportRow, to_csv, to_json, to_xlsx
from tests.conftest import good_values


def test_archive_path_matches_prd_example():
    values = {"issue_date": "2026-03-14", "vendor_name": "Toko Maju Jaya", "document_number": "NT-20260314-0231",
              "total": 1250000.0}  # fmt: skip
    assert (
        archive_path(values, "receipt", "pdf")
        == "Arsip/2026/03/Nota/2026-03-14_TokoMajuJaya_NT-20260314-0231_1250000.pdf"
    )


def test_archive_path_handles_missing_values_and_slashes():
    p = archive_path(
        {"vendor_name": "PT. Sinar/Abadi", "document_number": "INV/2026/03/0045", "total": None},
        "invoice",
        "jpg",
    )
    assert p.startswith("Arsip/TanpaTahun/00/Faktur/TanpaTanggal_PTSinarAbadi_INV-2026-03-0045_0")
    assert ".." not in p


def test_pattern_rejects_unknown_placeholders_and_traversal():
    with pytest.raises(ValueError):
        validate_pattern("{tanggal}_{password}")
    with pytest.raises(ValueError):
        validate_pattern("../{tanggal}")


def test_dedupe_paths():
    assert dedupe_paths(["a/x.pdf", "a/x.pdf", "a/y.pdf"]) == ["a/x.pdf", "a/x_2.pdf", "a/y.pdf"]


def rows():
    v = good_values()
    v["vendor_name"] = '=HYPERLINK("http://evil")'  # formula injection attempt
    return [ExportRow("doc_1", "receipt", "reviewed", v, "Arsip/x.pdf", "nota.jpg")]


def test_xlsx_is_a_readable_report():
    import zipfile
    from datetime import datetime

    data = to_xlsx(rows())
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Ringkasan", "Item"]
    ws = wb["Ringkasan"]
    assert ws["A1"].value == "Rekap Dokumen" and "1 dokumen" in ws["A2"].value
    headers = [c.value for c in ws[4]]
    assert headers[:4] == ["No", "Tanggal", "Jenis", "Vendor"] and "Mata uang" not in headers
    col = {h: i for i, h in enumerate(headers)}
    row = ws[5]
    assert isinstance(row[col["Tanggal"]].value, datetime)  # a real date, not text
    assert row[col["Tanggal"]].number_format == "dd/mm/yyyy"
    assert row[col["Total"]].number_format.startswith('"Rp"')
    assert row[col["Metode bayar"]].value == "Tunai"
    vendor = row[col["Vendor"]]
    assert vendor.value.startswith("=") and vendor.data_type == "s"  # formula injection stays inert text
    assert ws[6][col["Total"]].value == "=SUM(K5:K5)"  # totals row
    assert ws.column_dimensions["F"].width >= len("Nomor dokumen")  # headers are not cut off
    assert "Ringkasan" in ws.tables and ws.freeze_panes == "A5"
    items = wb["Item"]
    assert [c.value for c in items[5]][4] == "Beras 5 kg" and items.max_row == 7  # 2 items + totals
    sheet_xml = zipfile.ZipFile(io.BytesIO(data)).read("xl/worksheets/sheet1.xml").decode()
    assert 'numberStoredAsText="1"' in sheet_xml


def test_csv_and_json():
    text = to_csv(rows()).decode("utf-8-sig")
    parsed = list(csv.reader(io.StringIO(text), delimiter=";"))
    assert parsed[0][-1] == "Jumlah item" and parsed[1][-1] == "2"
    data = json.loads(to_json(rows()))
    assert data["documents"][0]["total"] == 270840.0
